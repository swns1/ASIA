"""
A saved AI summary as a Word (.docx) file.

The summary text uses the same light markup the page renders
(components/AIInsightPanel.jsx renderInterpretation): a line that is wholly
**BOLD** is a section heading, a line starting "* " or "- " is a bullet, and
**bold** / *italic* work inside a line. parse_blocks() follows those rules so
the file reads the way the screen does.
"""
import re
from io import BytesIO

from django.utils import timezone
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from enrollments.models import Enrollment
from enrollments.report_views import GRADING_PERIOD_LABELS
from school_forms.data import load_school

from .summaries import CONTEXT_TITLES

DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_HEADING = re.compile(r"^\*\*[^*]+\*\*$")
_BULLET = re.compile(r"^[*-]\s+")
_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")
_MUTED = RGBColor(0x6B, 0x5E, 0x5E)


def parse_inline(text):
    """[(text, bold, italic)] for one line."""
    runs = []
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            runs.append((part[2:-2], True, False))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            runs.append((part[1:-1], False, True))
        else:
            runs.append((part, False, False))
    return runs


def parse_blocks(text):
    """[(kind, payload)]: ("heading", str), ("bullet", runs), ("paragraph", runs).
    Blank lines separate paragraphs and are not blocks themselves."""
    blocks = []
    for raw in (text or "").split("\n"):
        line = raw.strip()
        if not line:
            continue
        if _HEADING.match(line):
            blocks.append(("heading", line.replace("**", "")))
        elif _BULLET.match(line):
            blocks.append(("bullet", parse_inline(_BULLET.sub("", line))))
        else:
            blocks.append(("paragraph", parse_inline(line)))
    return blocks


def _when(dt):
    """'October 14, 2026, 3:20 PM'. Built by hand: strftime's no-padding
    flag (%-d) does not exist on Windows, where the LAN deployment runs."""
    if not dt:
        return ""
    local = timezone.localtime(dt)
    hour = local.hour % 12 or 12
    return f"{local:%B} {local.day}, {local.year}, {hour}:{local:%M %p}"


def _student_name(student):
    middle = f" {student.middle_name[0]}." if student.middle_name else ""
    suffix = f" {student.suffix}" if student.suffix else ""
    return f"{student.last_name}, {student.first_name}{middle}{suffix}"


def describe_subject(summary):
    """(heading line, detail line) naming what the summary is about."""
    meta = summary.subject_meta or {}
    if "enrollment_id" in meta:
        enrollment = (
            Enrollment.objects.select_related("student").filter(pk=meta["enrollment_id"]).first()
        )
        if enrollment is None:
            return "Student record no longer on file", ""
        detail = " · ".join(p for p in (
            enrollment.grade_level, enrollment.section, f"SY {enrollment.school_year}",
        ) if p)
        return _student_name(enrollment.student), detail
    period = meta.get("grading_period") or ""
    parts = [
        meta.get("grade_level") or "All grade levels",
        GRADING_PERIOD_LABELS.get(period, period.replace("_", " ").title()),
        meta.get("subject") or "",
        f"SY {meta['school_year']}" if meta.get("school_year") else "",
    ]
    return "Student grouping", " · ".join(p for p in parts if p)


def file_stem(summary, subject_heading):
    """'AI-Grade-Interpretation_Dela-Cruz-Juan-M': ASCII letters, digits, dashes."""
    title = CONTEXT_TITLES.get(summary.context_type, "AI Summary")
    cleaned = [re.sub(r"[^A-Za-z0-9-]+", "-", p).strip("-") for p in (title, subject_heading)]
    return "_".join(p for p in cleaned if p)[:120] or "AI-Summary"


def _add_runs(paragraph, runs):
    for text, bold, italic in runs:
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic


def _muted(paragraph, text, size=9, italic=False):
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = _MUTED
    run.italic = italic
    return run


def build_docx(summary, *, text, byline):
    """The file's bytes. `text` is the version this reader may see (the
    owner's current one, or the published copy); `byline` says whose."""
    school = load_school()
    heading, detail = describe_subject(summary)

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    if school.name:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(school.name).bold = True
        if school.address:
            p.add_run("\n")
            _muted(p, school.address)

    doc.add_heading(CONTEXT_TITLES.get(summary.context_type, "AI Summary"), level=1)
    p = doc.add_paragraph()
    p.add_run(heading).bold = True
    if detail:
        p.add_run("\n" + detail)
    _muted(doc.add_paragraph(), byline)

    for kind, payload in parse_blocks(text):
        if kind == "heading":
            doc.add_heading(payload.title(), level=2)
        elif kind == "bullet":
            _add_runs(doc.add_paragraph(style="List Bullet"), payload)
        else:
            _add_runs(doc.add_paragraph(), payload)

    _muted(doc.add_paragraph(), "AI-generated · For guidance purposes only", italic=True)

    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def byline_for(summary, *, is_owner, owner_name):
    """Who wrote this version and when, for the line under the heading."""
    provider = {"gemini": "Gemini", "groq": "Groq"}.get(summary.provider, "AI")
    if not is_owner:
        return f"Shared by {owner_name} on {_when(summary.shared_at)}"
    line = f"Generated by {provider} on {_when(summary.generated_at)}"
    if summary.edited_text:
        line += f" · Edited by {owner_name} on {_when(summary.edited_at)}"
    return line
