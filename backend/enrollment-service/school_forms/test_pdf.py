"""
The school forms as PDF: numbers shown the way Excel shows them, text set
and clipped the way it prints, pages cut where Excel cuts them.

Most checks draw onto a Recorder, a stand-in pen that keeps what it was
asked to draw; the real PDF is checked end to end at the bottom. The page
counts there are Excel's own, read off its printouts of the same files.
"""
import re
from contextlib import contextmanager
from datetime import date
from types import SimpleNamespace

import openpyxl
import pytest
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.pagebreak import Break

from school_forms import pdf, sf1
from school_forms.test_sf1 import learner, register
from school_forms.test_sf10 import GRADE_1, GRADE_5, HUMSS, KINDER, build as build_sf10

THIN = Side(style="thin")


class Recorder:
    """A pen that remembers what it was asked to draw, page by page."""

    def __init__(self):
        self.pages = []
        self.clips = []

    def page(self, width, height):
        self.pages.append(SimpleNamespace(size=(width, height), texts=[], lines=[], fills=[], images=[]))

    def text(self, x, y, text, style, angle=0):
        clip = self.clips[-1] if self.clips else None
        self.pages[-1].texts.append(SimpleNamespace(x=x, y=y, text=text, style=style, angle=angle, clip=clip))

    def line(self, x1, y1, x2, y2, width, color, dash=None):
        self.pages[-1].lines.append(SimpleNamespace(x1=x1, y1=y1, x2=x2, y2=y2, width=width, dash=dash))

    def fill(self, x, y, w, h, color):
        self.pages[-1].fills.append((x, y, w, h, color))

    def image(self, data, x, y, w, h):
        self.pages[-1].images.append((x, y, w, h))

    @contextmanager
    def clip(self, x, y, w, h):
        self.clips.append((x, y, w, h))
        try:
            yield
        finally:
            self.clips.pop()


def drawn(ws, stretch=pdf.AS_DRAWN):
    pen = Recorder()
    pdf.draw(ws, pen, stretch)
    return pen


def sheet():
    ws = openpyxl.Workbook().active
    ws.page_setup.paperSize = 1
    for side in ("left", "right", "top", "bottom"):
        setattr(ws.page_margins, side, 0.5)
    return ws


def text(page, value):
    return next(t for t in page.texts if t.text == value)


# ── Numbers and dates ───────────────────────────────────────────────────────

@pytest.mark.parametrize("value, fmt, shown", [
    (94, "General", "94"),
    (94.5, "General", "94.5"),
    (94.5, "0", "95"),                      # half up, as Excel rounds
    (94, "0.00", "94.00"),
    (1234567, "#,##0", "1,234,567"),
    (0.256, "0%", "26%"),
    (0.256, "0.0%", "25.6%"),
    (42, "00000", "00042"),
    (660, '0" hours"', "660 hours"),
    (12, "#.###", "12."),                   # DepEd's age format keeps the point
    (-5, "0.00", "-5.00"),
    (date(2013, 11, 5), "m/d/yy;@", "11/5/13"),
    (date(2013, 11, 5), "[$-10464]m/d/yyyy;@", "11/5/2013"),
    (date(2026, 9, 1), "mmm-yy", "Sep-26"),
    (date(2026, 9, 1), "mm-dd-yy", "09-01-26"),
    (date(2026, 3, 31), "mmmm d, yyyy", "March 31, 2026"),
    ("as typed", "0.00", "as typed"),
    (None, "General", ""),
])
def test_values_show_as_excel_shows_them(value, fmt, shown):
    assert pdf.display(value, fmt) == shown


def test_text_keeps_to_the_built_in_fonts():
    # Accents the fonts have stay; marks drawn as shapes stay; the rest lose their accent or become ?.
    assert pdf.printable("Peña – ◤ ✓ Lăng") == "Peña – ◤ ? Lang"


# ── Setting text ────────────────────────────────────────────────────────────

def test_wrapping_breaks_at_spaces_and_keeps_an_opening_indent():
    style = pdf.style_of("Times New Roman", 10, False, False, None, None, None, 1.0)
    lines = pdf.lines_of([("    I hereby certify that this is a true record of the learner.", style)], limit=150)
    assert len(lines) == 2
    assert lines[0].runs[0][0].startswith("    I hereby")          # the paragraph's indent survives
    assert not lines[1].runs[0][0].startswith(" ")                  # a wrapped line doesn't
    assert all(line.width <= 150 for line in lines)


def test_a_word_longer_than_the_line_is_broken():
    style = pdf.style_of("Arial", 10, False, False, None, None, None, 1.0)
    lines = pdf.lines_of([("Supercalifragilisticexpialidocious", style)], limit=40)
    assert len(lines) > 1 and all(line.width <= 40 for line in lines)


def test_alignment_places_text_in_its_cell():
    ws = sheet()
    ws["A1"], ws["A2"], ws["A3"] = "left", "middle", "right"
    ws["A2"].alignment = Alignment(horizontal="center")
    ws["A3"].alignment = Alignment(horizontal="right")
    page = drawn(ws).pages[0]
    left, middle, right = text(page, "left"), text(page, "middle"), text(page, "right")
    cell_left, cell_width = 36, 64 * 0.75
    assert left.x == pytest.approx(cell_left + pdf.PAD)
    assert middle.x + middle.style.width("middle") / 2 == pytest.approx(cell_left + cell_width / 2)
    assert right.x + right.style.width("right") == pytest.approx(cell_left + cell_width - pdf.PAD)


def test_a_long_label_runs_on_over_empty_cells_and_stops_at_a_filled_one():
    ws = sheet()
    ws["A1"] = "A label much longer than its cell, and then some more"
    ws["D1"] = "next"
    page = drawn(ws).pages[0]
    label = text(page, ws["A1"].value)
    column = 64 * 0.75
    x, _, width, _ = label.clip
    assert x == pytest.approx(36) and width == pytest.approx(3 * column)   # A to C, not into D


def test_shrink_to_fit_makes_the_text_smaller_not_clipped():
    ws = sheet()
    ws["A1"] = "Senior High School Department Head"
    ws["A1"].alignment = Alignment(shrink_to_fit=True)
    label = text(drawn(ws).pages[0], "Senior High School Department Head")
    assert label.style.size < 11 and label.style.width(label.text) <= 64 * 0.75 - 2 * pdf.PAD + 0.01


def test_rotated_text_is_turned():
    ws = sheet()
    ws["A1"], ws["B1"] = "Up", "Down"
    ws["A1"].alignment = Alignment(textRotation=90)
    ws["B1"].alignment = Alignment(textRotation=180)        # Excel's 91-180 turn the other way
    page = drawn(ws).pages[0]
    assert text(page, "Up").angle == 90 and text(page, "Down").angle == -90


def test_rich_text_keeps_each_runs_font():
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont

    ws = sheet()
    ws["A1"] = CellRichText("21", TextBlock(InlineFont(vertAlign="superscript"), "st"), " Century")
    page = drawn(ws).pages[0]
    base, raised = text(page, "21"), text(page, "st")
    assert raised.style.size < base.style.size and raised.y < base.y


# ── Borders ─────────────────────────────────────────────────────────────────

def test_a_merged_range_has_borders_only_round_its_edge():
    ws = sheet()
    ws["A1"] = "Name"
    ws["A1"].border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws.merge_cells("A1:C1")
    lines = drawn(ws).pages[0].lines
    verticals = {round(line.x1, 2) for line in lines if line.x1 == line.x2}
    column = 64 * 0.75
    assert verticals == {36.0, round(36 + 3 * column, 2)}       # no line inside the range


def test_diagonal_borders_are_drawn():
    ws = sheet()
    ws["A1"].border = Border(diagonal=Side(style="dotted"), diagonalUp=True)
    lines = drawn(ws).pages[0].lines
    (diagonal,) = [line for line in lines if line.x1 != line.x2 and line.y1 != line.y2]
    assert diagonal.y1 > diagonal.y2 and diagonal.dash          # bottom left to top right, dotted


# ── Pages ───────────────────────────────────────────────────────────────────

def test_manual_breaks_start_new_pages():
    ws = sheet()
    for r in range(1, 11):
        ws.cell(r, 1, f"row {r}")
    ws.row_breaks.append(Break(id=4))                           # after row 4
    pages = drawn(ws).pages
    assert len(pages) == 2 and text(pages[1], "row 5").y < text(pages[1], "row 6").y
    assert [t.text for t in pages[0].texts] == ["row 1", "row 2", "row 3", "row 4"]


def test_title_rows_repeat_on_every_page():
    ws = sheet()
    ws["A1"] = "LEARNER'S NAME"
    for r in range(2, 120):
        ws.cell(r, 1, f"learner {r}")
    ws.print_title_rows = "1:1"
    pages = drawn(ws).pages
    assert len(pages) > 1
    assert all(page.texts[0].text == "LEARNER'S NAME" for page in pages)


def test_pages_with_nothing_to_print_are_left_out():
    ws = sheet()
    ws["A1"] = "printed"
    ws["Z1"].font = Font(bold=True)                             # formatting alone prints nothing
    ws.col_breaks.append(Break(id=1))
    assert len(drawn(ws).pages) == 1


def test_fitting_one_page_wide_shrinks_but_never_enlarges():
    narrow, wide = sheet(), sheet()
    for ws in (narrow, wide):
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    narrow["A1"] = "x"
    for c in range(1, 40):
        wide.cell(1, c, "x")
    assert pdf.Layout(narrow).scale == 1.0
    layout = pdf.Layout(wide)
    assert layout.scale < 1 and (layout.x[40] - layout.x[1]) * layout.sx == pytest.approx(layout.room[0])


def test_excels_printout_stretches_the_grid_but_not_the_text():
    ws = sheet()
    ws["C1"] = "here"
    as_drawn, printed = text(drawn(ws).pages[0], "here"), text(drawn(ws, pdf.EXCEL_PRINT).pages[0], "here")
    assert printed.x - 36 == pytest.approx((as_drawn.x - 36) * 1.075)
    assert printed.style.size == as_drawn.style.size


# ── The forms, end to end ───────────────────────────────────────────────────

def pages_in(data):
    return len(re.findall(rb"/Type /Page\b(?!s)", data))


def test_an_sf10_prints_as_drawn_on_long_bond():
    ws, _ = build_sf10("elementary", [GRADE_1, GRADE_5], prior=KINDER)
    pen = drawn(ws)
    assert [page.size for page in pen.pages] == [(612, 936)] * len(pen.pages)
    assert len(pen.pages[0].images) == 2                                       # the crest and the photo box
    footer = "".join(t.text for t in pen.pages[-1].texts if t.y > 900)
    assert footer == f"Page {len(pen.pages)} of {len(pen.pages)}"
    data = pdf.render(ws)
    assert data.startswith(b"%PDF-") and pages_in(data) == len(pen.pages)


def test_the_three_sf10_signatures_print_side_by_side():
    ws, _ = build_sf10("senior_highschool", HUMSS)
    page = drawn(ws).pages[-1]
    names = [text(page, n) for n in ("Regina T. Molidor, LPT", "Erwin V. Bustamante, M. Div.",
                                     "Maria Carolina A. Laguerta, LPT")]
    assert len({round(n.y, 2) for n in names}) == 1 and names[0].x < names[1].x < names[2].x


def test_sf1_paginates_as_excel_prints_it():
    # Excel prints this class of 26 on two Legal pages, the header rows repeated (read off its own printout).
    learners = [learner(i + 1, "male") for i in range(12)] + [learner(100 + i, "female") for i in range(14)]
    wb, _ = sf1.build(register(learners))
    pen = drawn(wb.active, pdf.EXCEL_PRINT)
    assert [page.size for page in pen.pages] == [(1008, 612)] * 2
    assert any(t.text.startswith("LRN") for t in pen.pages[1].texts)
    data = pdf.render(wb.active, stretch=pdf.EXCEL_PRINT)
    assert pages_in(data) == 2
