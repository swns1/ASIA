"""
A school form as a PDF: the sheet its Excel download builds, laid out the
way Excel prints it.

The server needs nothing but reportlab. The sheet's print setup (paper,
orientation, margins, scale, rows repeated on every page, page breaks) and
its cells (merged ranges, borders, fills, fonts, alignment, wrapped, shrunk
and rotated text, rich text, number formats, pictures, the footer) are read
here and drawn as Excel prints them. That covers what DepEd's SF1/SF2 files
and the drawn SF10 use; it is not every feature Excel has.

Excel prints a sheet's grid wider and shorter than it shows it: measured
off its own PDFs, columns come out 7.5% wider and rows 3.5% shorter, while
text and pictures keep their size. DepEd laid its SF1/SF2 files out for
that printout (what fits on a page, where its logos sit beside the titles),
so their PDFs stretch the grid the same way (EXCEL_PRINT). The SF10 is
drawn to the school's Word templates in exact inches and is printed as
drawn (AS_DRAWN).

The fonts are the PDF's built-in ones, so no font file ships with the code:
- Times for Times New Roman and the other serifs. It has Times New Roman's
  widths, so an SF10 sets line for line as it does in Excel.
- Helvetica for Arial and the other sans serifs, set narrower for Arial
  Narrow (Arial at 82% of the width) and Calibri.
"""
import re
import unicodedata
from collections import defaultdict
from contextlib import contextmanager
from dataclasses import dataclass, replace
from datetime import date, datetime, time
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from io import BytesIO
from math import ceil, cos, radians, sin

from openpyxl.cell.cell import MergedCell
from openpyxl.cell.rich_text import CellRichText
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen.canvas import Canvas

PDF_CONTENT_TYPE = "application/pdf"

# How far the printed grid is from the one on screen, (across, down).
EXCEL_PRINT = (1.075, 0.965)
AS_DRAWN = (1.0, 1.0)

INCH = 72
MM = 72 / 25.4
EMU_PER_POINT = 12700
# Excel's paper codes. Above 118 a code is the printer's own; DepEd's senior
# high forms carry 258, the long bond (8.5 x 13 in) their printers call it.
PAPER = {
    1: (8.5 * INCH, 11 * INCH),        # Letter
    5: (8.5 * INCH, 14 * INCH),        # Legal
    8: (297 * MM, 420 * MM),           # A3
    9: (210 * MM, 297 * MM),           # A4
    11: (148 * MM, 210 * MM),          # A5
    14: (8.5 * INCH, 13 * INCH),       # Folio, the Philippines' long bond
}
LONG_BOND = PAPER[14]

PAD = 1.5                    # points Excel keeps between a cell's edge and its text
PAD_V = 1.0
INDENT = 9.0                 # points per indent level

# Built-in PDF faces by family, then (bold, italic).
FACES = {
    "Times": {(False, False): "Times-Roman", (True, False): "Times-Bold",
              (False, True): "Times-Italic", (True, True): "Times-BoldItalic"},
    "Helvetica": {(False, False): "Helvetica", (True, False): "Helvetica-Bold",
                  (False, True): "Helvetica-Oblique", (True, True): "Helvetica-BoldOblique"},
    "Courier": {(False, False): "Courier", (True, False): "Courier-Bold",
                (False, True): "Courier-Oblique", (True, True): "Courier-BoldOblique"},
}
SERIFS = {"times new roman", "times", "century schoolbook", "century", "cambria", "georgia", "book antiqua",
          "garamond", "bookman old style", "palatino linotype", "constantia", "serif"}
MONOSPACED = {"courier new", "courier", "consolas", "lucida console"}
NARROW = {"arial narrow": 82, "calibri": 90}       # percent of the face's width
# Ascent, descent and line spacing in ems, as Windows sets Times New Roman,
# Arial and Courier New in a cell.
METRICS = {"Times": (0.891, 0.216, 1.15), "Helvetica": (0.905, 0.212, 1.15), "Courier": (0.833, 0.300, 1.13)}

# Marks the built-in fonts lack, drawn as shapes: corners of a unit square,
# top down. SF2 marks a learner tardy with ◤.
SHAPES = {
    "◤": ((0, 0), (1, 0), (0, 1)), "◥": ((0, 0), (1, 0), (1, 1)),
    "◣": ((0, 0), (1, 1), (0, 1)), "◢": ((1, 0), (1, 1), (0, 1)),
    "■": ((0, 0), (1, 0), (1, 1), (0, 1)),
}
SHAPE_EM = 0.75              # a shape's side, in ems

# Office's default theme colours, for a fill or font that names one.
THEME = ("FFFFFF", "000000", "E7E6E6", "44546A", "4472C4", "ED7D31", "A5A5A5", "FFC000", "5B9BD5", "70AD47")
BLACK = (0.0, 0.0, 0.0)
WHITE = (1.0, 1.0, 1.0)

# Line widths in points, and dash patterns, for Excel's border styles.
BORDERS = {
    "hair": (0.25, None), "thin": (0.5, None), "medium": (1.0, None), "thick": (1.5, None),
    "dotted": (0.5, (0.5, 1.5)), "dashed": (0.5, (3, 1.5)), "dashDot": (0.5, (3, 1.5, 0.5, 1.5)),
    "dashDotDot": (0.5, (3, 1.5, 0.5, 1.5, 0.5, 1.5)), "mediumDashed": (1.0, (4, 2)),
    "mediumDashDot": (1.0, (4, 2, 1, 2)), "mediumDashDotDot": (1.0, (4, 2, 1, 2, 1, 2)),
    "slantDashDot": (1.0, (4, 2, 1, 2)), "double": (0.5, None),
}


def render(ws, title=None, stretch=AS_DRAWN):
    """The worksheet `ws` as PDF bytes, one PDF page per printed page."""
    pen = Pen(title or ws.title)
    draw(ws, pen, stretch)
    return pen.finish()


def draw(ws, pen, stretch=AS_DRAWN):
    """Lay `ws` out and draw it with `pen` (a Pen, or a stand-in that records)."""
    layout = Layout(ws, stretch)
    pages = layout.pages()
    for number, page in enumerate(pages, 1):
        pen.page(*layout.paper)
        for segment in page.segments:
            layout.paint(pen, page, segment)
        layout.header_footer(pen, number, len(pages))


# ─── Numbers as Excel shows them ────────────────────────────────────────────

def display(value, fmt="General"):
    """`value` as Excel shows it under number format `fmt` (the formats DepEd's forms use)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, str):
        return value
    sections = _sections(fmt or "General")
    if isinstance(value, (datetime, date, time)):
        return _date_text(value, sections[0])
    if isinstance(value, (int, float, Decimal)):
        return _number_text(value, sections)
    return str(value)


def _sections(fmt):
    """A format's ;-separated sections, ignoring ; inside quotes."""
    out, current, quoted = [], "", False
    for ch in fmt:
        if ch == '"':
            quoted = not quoted
        if ch == ";" and not quoted:
            out.append(current)
            current = ""
        else:
            current += ch
    out.append(current)
    return out


def _literal(section):
    """A section without its [locale] and [colour] codes."""
    return re.sub(r"\[[^\]]*\]", "", section)


def _general(number):
    if number == number.to_integral_value():
        return str(int(number))
    text = format(float(number), ".10g")
    return text if "e" not in text else format(float(number), ".5E")


def _number_text(value, sections):
    try:
        number = Decimal(str(value))
    except InvalidOperation:
        return str(value)
    sign = ""
    if number < 0 and len(sections) > 1:
        section, number = sections[1], -number
    elif number == 0 and len(sections) > 2:
        section = sections[2]
    else:
        section = sections[0]
        if number < 0:
            sign, number = "-", -number
    section = _literal(section)
    if not section.strip() or section.strip().lower() == "general":
        return sign + _general(number)

    # Literal text around the digits: "quoted", \escaped, _x (a space), *x (a fill, dropped).
    parts, digits = [], None
    for token in re.findall(r'"[^"]*"|\\.|_.|\*.|[0#?,.]+|.', section):
        if re.fullmatch(r"[0#?,.]+", token) and digits is None and re.search(r"[0#?]", token):
            digits = token
            parts.append(None)
        elif token.startswith('"'):
            parts.append(token[1:-1])
        elif token.startswith("\\"):
            parts.append(token[1:])
        elif token.startswith("_"):
            parts.append(" ")
        elif token.startswith("*"):
            continue
        else:
            parts.append(token)
    if digits is None:                       # no number in it: Excel prints the literal
        return sign + "".join(p for p in parts if p)
    if "%" in "".join(p for p in parts if p):
        number *= 100

    whole, _, fraction = digits.partition(".")
    places = len(re.findall(r"[0#?]", fraction))
    required = fraction.count("0")
    number = number.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
    int_text, _, frac_text = f"{number:f}".partition(".")
    frac_text = frac_text.ljust(places, "0")[:places]
    while len(frac_text) > required and frac_text.endswith("0"):
        frac_text = frac_text[:-1]
    int_text = int_text.lstrip("0").rjust(whole.count("0"), "0")
    if "," in whole and int_text:
        int_text = f"{int(int_text):,}"
    text = int_text + ("." + frac_text if "." in digits else "")
    return sign + "".join(text if p is None else p for p in parts)


def _date_text(value, section):
    section = _literal(section).replace("@", "")
    if isinstance(value, time):
        value = datetime.combine(date(1899, 12, 30), value)
    if not isinstance(value, datetime):
        value = datetime(value.year, value.month, value.day)
    tokens = re.findall(r'"[^"]*"|\\.|yyyy|yy|mmmmm|mmmm|mmm|mm|m|dddd|ddd|dd|d|hh|h|ss|s|am/pm|a/p|.',
                        section, flags=re.IGNORECASE)
    twelve_hour = any(t.lower() in ("am/pm", "a/p") for t in tokens)
    out = []
    for i, token in enumerate(tokens):
        low = token.lower()
        # m is minutes right after an hour or right before seconds.
        minutes = low in ("m", "mm") and (
            any(t.lower() in ("h", "hh") for t in tokens[max(0, i - 2):i])
            or any(t.lower() in ("s", "ss") for t in tokens[i + 1:i + 3]))
        if token.startswith('"'):
            out.append(token[1:-1])
        elif token.startswith("\\"):
            out.append(token[1:])
        elif low == "yyyy":
            out.append(f"{value.year:04d}")
        elif low == "yy":
            out.append(f"{value.year % 100:02d}")
        elif minutes:
            out.append(f"{value.minute:02d}" if low == "mm" else str(value.minute))
        elif low == "mmmmm":
            out.append(value.strftime("%B")[0])
        elif low == "mmmm":
            out.append(value.strftime("%B"))
        elif low == "mmm":
            out.append(value.strftime("%b"))
        elif low == "mm":
            out.append(f"{value.month:02d}")
        elif low == "m":
            out.append(str(value.month))
        elif low == "dddd":
            out.append(value.strftime("%A"))
        elif low == "ddd":
            out.append(value.strftime("%a"))
        elif low == "dd":
            out.append(f"{value.day:02d}")
        elif low == "d":
            out.append(str(value.day))
        elif low in ("hh", "h"):
            hour = value.hour % 12 or 12 if twelve_hour else value.hour
            out.append(f"{hour:02d}" if low == "hh" else str(hour))
        elif low in ("ss", "s"):
            out.append(f"{value.second:02d}" if low == "ss" else str(value.second))
        elif low == "am/pm":
            out.append("AM" if value.hour < 12 else "PM")
        elif low == "a/p":
            out.append("A" if value.hour < 12 else "P")
        else:
            out.append(token)
    return "".join(out)


# ─── Fonts, colours, text runs ──────────────────────────────────────────────

@dataclass(frozen=True)
class Style:
    """How a run of text is set: a built-in PDF face, already scaled to the page."""

    font: str
    size: float
    hscale: float = 100.0
    color: tuple = BLACK
    underline: bool = False
    rise: float = 0.0            # above the baseline (superscript)
    ascent: float = 0.891        # in ems
    descent: float = 0.216
    spacing: float = 1.15

    def width(self, text):
        shapes = sum(text.count(mark) for mark in SHAPES)
        if shapes:
            text = "".join(ch for ch in text if ch not in SHAPES)
        return stringWidth(text, self.font, self.size) * self.hscale / 100 + shapes * SHAPE_EM * self.size

    def scaled(self, factor):
        return replace(self, size=self.size * factor, rise=self.rise * factor)


def family(name):
    key = (name or "").strip().lower()
    if key in MONOSPACED:
        return "Courier"
    if key in SERIFS or key.startswith("times"):
        return "Times"
    return "Helvetica"


def rgb(color, default=None):
    """An openpyxl colour as (r, g, b) in 0..1, or `default` when it has none."""
    if color is None:
        return default
    value = None
    kind = getattr(color, "type", None)
    if kind == "rgb" and isinstance(color.rgb, str):
        value = color.rgb
    elif kind == "indexed":
        index = color.indexed
        value = COLOR_INDEX[index] if index is not None and index < len(COLOR_INDEX) else None
        if index == 64:                       # the system's text colour
            value = "000000"
    elif kind == "theme" and color.theme is not None and color.theme < len(THEME):
        value = THEME[color.theme]
    if not value or len(value) < 6:
        return default
    value = value[-6:]
    try:
        return tuple(int(value[i:i + 2], 16) / 255 for i in (0, 2, 4))
    except ValueError:
        return default


def style_of(name, size, bold, italic, underline, vert_align, color, scale):
    face = family(name)
    ascent, descent, spacing = METRICS[face]
    size = float(size or 11) * scale
    rise = 0.0
    if vert_align == "superscript":
        rise, size = size * 0.33, size * 0.67
    elif vert_align == "subscript":
        rise, size = -size * 0.12, size * 0.67
    return Style(
        font=FACES[face][bool(bold), bool(italic)], size=size,
        hscale=NARROW.get((name or "").strip().lower(), 100), color=rgb(color, BLACK),
        underline=bool(underline) and underline != "none", rise=rise,
        ascent=ascent, descent=descent, spacing=spacing,
    )


def printable(text):
    """
    `text` in the built-in fonts' character set (and the shapes drawn for
    them): other letters lose their accent, anything else becomes ?.
    """
    out = []
    for ch in text:
        if ch in SHAPES:
            out.append(ch)
            continue
        try:
            ch.encode("cp1252")
            out.append(ch)
        except UnicodeEncodeError:
            base = unicodedata.normalize("NFKD", ch).encode("cp1252", "ignore").decode("cp1252")
            out.append(base or "?")
    return "".join(out)


def runs_of(cell, scale):
    """The cell's text as [(text, Style)]: one run, or one per rich-text part."""
    f = cell.font
    base = style_of(f.name, f.sz, f.b, f.i, f.u, f.vertAlign, f.color, scale)
    value = cell.value
    if isinstance(value, CellRichText):
        runs = []
        for part in value:
            if isinstance(part, str):
                runs.append((part, base))
                continue
            inline = part.font
            runs.append((part.text, style_of(
                inline.rFont or f.name, inline.sz or f.sz,
                f.b if inline.b is None else inline.b, f.i if inline.i is None else inline.i,
                f.u if inline.u is None else inline.u, inline.vertAlign or f.vertAlign,
                inline.color or f.color, scale)))
        return [(printable(t), s) for t, s in runs if t]
    text = display(value, cell.number_format)
    return [(printable(text), base)] if text else []


@dataclass
class Line:
    runs: list                   # [(text, Style)]
    width: float
    ascent: float
    height: float


def _line(runs, fallback):
    runs = _merge(runs)
    styles = [s for _, s in runs] or [fallback]
    return Line(
        runs=runs, width=sum(s.width(t) for t, s in runs),
        ascent=max(s.ascent * s.size for s in styles),
        height=max(s.spacing * s.size for s in styles),
    )


def _merge(runs):
    out = []
    for text, style in runs:
        if out and out[-1][1] == style:
            out[-1] = (out[-1][0] + text, style)
        else:
            out.append((text, style))
    return out


def _rstrip(runs):
    runs = list(runs)
    while runs and not runs[-1][0].strip():
        runs.pop()
    if runs:
        runs[-1] = (runs[-1][0].rstrip(" "), runs[-1][1])
    return runs


def lines_of(runs, limit=None):
    """
    The runs set as lines: broken at newlines, and when `limit` is given at
    the spaces that keep each line within it (a word longer than a line is
    broken where it must, as Excel does). Without a limit, newlines are
    spaces: a cell that doesn't wrap prints on one line.
    """
    fallback = runs[0][1] if runs else Style("Helvetica", 11)
    if limit is None:
        return [_line([(t.replace("\n", " "), s) for t, s in runs], fallback)]
    tokens = [(tok, style) for text, style in runs for tok in re.findall(r"\n| +|[^ \n]+", text)]
    lines, current, width = [], [], 0.0
    hard_start = True                # spaces opening the text or a line after a newline indent it
    for tok, style in tokens:
        if tok == "\n":
            lines.append(_line(_rstrip(current), fallback))
            current, width, hard_start = [], 0.0, True
            continue
        w = style.width(tok)
        if tok.isspace():
            if current or hard_start:
                current.append((tok, style))
                width += w
            continue
        hard_start = False
        if current and width + w > limit and any(t.strip() for t, _ in current):
            lines.append(_line(_rstrip(current), fallback))
            current, width = [], 0.0
        while not current and w > limit and len(tok) > 1:
            cut = len(tok) - 1
            while cut > 1 and style.width(tok[:cut]) > limit:
                cut -= 1
            lines.append(_line([(tok[:cut], style)], fallback))
            tok = tok[cut:]
            w = style.width(tok)
        current.append((tok, style))
        width += w
    lines.append(_line(_rstrip(current), fallback))
    return lines


# ─── Laying the sheet on pages ──────────────────────────────────────────────

@dataclass
class Page:
    cols: tuple                  # first, last column printed
    segments: list               # (first row, last row, top in points); the title rows come first


def _max_digit_width(wb):
    """Pixels per digit of the workbook's default font: Excel's unit of column width."""
    try:
        font = wb._fonts[0]
        name, size = (font.name or "Calibri").lower(), float(font.sz or 11)
    except (AttributeError, IndexError, TypeError):
        name, size = "calibri", 11.0
    em = {"calibri": 0.5068, "arial": 0.556, "times new roman": 0.5, "arial narrow": 0.456,
          "cambria": 0.5562, "verdana": 0.6362, "tahoma": 0.5459}.get(name, 0.5068)
    return max(1, round(em * size * 96 / 72))


def _shows(cell):
    """Whether a cell puts anything on paper: a value, a border or a coloured fill."""
    if cell.value not in (None, ""):
        return True
    border = cell.border
    if border is not None:
        if any(getattr(border, side) is not None and getattr(border, side).style
               for side in ("left", "right", "top", "bottom")):
            return True
        if (border.diagonalUp or border.diagonalDown) and border.diagonal is not None and border.diagonal.style:
            return True
    fill = cell.fill
    return fill is not None and fill.fill_type == "solid" and rgb(fill.fgColor, WHITE) != WHITE


def _title_rows(ws):
    numbers = [int(n) for n in re.findall(r"\d+", ws.print_title_rows or "")]
    return (min(numbers), max(numbers)) if numbers else None


@dataclass
class Picture:
    data: bytes
    start: tuple                 # (column, points in, row, points down): where its top left is
    end: tuple = None            # the same for its bottom right, when it stretches with the cells
    size: tuple = None           # (width, height) in points, when it keeps its own size


class Layout:
    """One worksheet's geometry and print setup, and the painting of its pages."""

    def __init__(self, ws, stretch=AS_DRAWN):
        self.ws = ws
        self.rows = defaultdict(list)
        self.showing = set()
        for (r, c), cell in ws._cells.items():
            self.rows[r].append(cell)
            if _shows(cell):
                self.showing.add((r, c))
        self.pictures = list(self._pictures())
        self.min_row, self.min_col, self.max_row, self.max_col = self._area()

        corners = [p.start for p in self.pictures] + [p.end for p in self.pictures if p.end]
        last_col = max(self.max_col, ws.max_column, *(corner[0] for corner in corners)) + 2
        last_row = max(self.max_row, ws.max_row, *(corner[2] for corner in corners)) + 2
        widths = self._column_widths(_max_digit_width(ws.parent), last_col)
        self.x = [0.0] * (last_col + 2)
        for c in range(1, last_col + 1):
            self.x[c + 1] = self.x[c] + widths.get(c, 0.0)
        default_height = ws.sheet_format.defaultRowHeight or 15
        self.y = [0.0] * (last_row + 2)
        for r in range(1, last_row + 1):
            dim = ws.row_dimensions.get(r)
            if dim is not None and dim.hidden:
                height = 0.0
            elif dim is not None and dim.height is not None:
                height = float(dim.height)
            else:
                height = float(default_height)
            self.y[r + 1] = self.y[r] + height

        setup = ws.page_setup
        paper = PAPER.get(int(setup.paperSize), LONG_BOND) if setup.paperSize else PAPER[1]
        self.paper = (paper[1], paper[0]) if setup.orientation == "landscape" else paper
        m = ws.page_margins
        self.margins = {side: float(getattr(m, side) if getattr(m, side) is not None else default) * INCH
                        for side, default in (("left", 0.7), ("right", 0.7), ("top", 0.75), ("bottom", 0.75),
                                              ("header", 0.3), ("footer", 0.3))}
        self.room = (self.paper[0] - self.margins["left"] - self.margins["right"],
                     self.paper[1] - self.margins["top"] - self.margins["bottom"])
        self.titles = _title_rows(ws)
        self.scale = self._scale(stretch)
        self.sx, self.sy = self.scale * stretch[0], self.scale * stretch[1]

        self.merged = {}                         # top-left (row, col) -> (min_row, min_col, max_row, max_col)
        self.owner = {}                          # every cell of a merged range -> its bounds
        for rng in ws.merged_cells.ranges:
            bounds = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)
            self.merged[rng.min_row, rng.min_col] = bounds
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    self.owner[r, c] = bounds

    def _area(self):
        """The print area, or (none being set) every cell and picture that shows, as Excel prints."""
        area = self.ws.print_area
        if area:
            if isinstance(area, (list, tuple)):
                area = area[0]
            ref = area.split(",")[0].split("!")[-1].replace("$", "")
            min_col, min_row, max_col, max_row = range_boundaries(ref)
            return min_row, min_col, max_row, max_col
        rows = [r for r, _ in self.showing] + [p.start[2] for p in self.pictures]
        cols = [c for _, c in self.showing] + [p.start[0] for p in self.pictures]
        if not rows:
            return 1, 1, 1, 1
        return min(rows), min(cols), max(rows), max(cols)

    def _column_widths(self, mdw, last_col):
        """Points per column. Excel stores a width in characters of the default font's digits."""
        def points(chars):
            pixels = int(((256 * chars + int(128 / mdw)) / 256) * mdw)
            return pixels * 0.75

        fmt = self.ws.sheet_format
        if fmt.defaultColWidth is not None:
            default = points(fmt.defaultColWidth)
        else:
            # Excel rounds a default column up to a whole 8 pixels: 64 for Calibri 11.
            default = ceil(((fmt.baseColWidth or 8) * mdw + 5) / 8) * 8 * 0.75
        widths = {c: default for c in range(1, last_col + 1)}
        for key, dim in self.ws.column_dimensions.items():
            first = dim.min or coordinate_to_tuple(f"{key}1")[1]
            last = dim.max or first
            for c in range(first, min(last, last_col) + 1):
                if dim.hidden:
                    widths[c] = 0.0
                elif dim.width:
                    widths[c] = points(dim.width)
        return widths

    def _scale(self, stretch):
        setup, props = self.ws.page_setup, self.ws.sheet_properties.pageSetUpPr
        if props is not None and props.fitToPage:
            wide = 1 if setup.fitToWidth is None else int(setup.fitToWidth)
            tall = 1 if setup.fitToHeight is None else int(setup.fitToHeight)
            width = (self.x[self.max_col + 1] - self.x[self.min_col]) * stretch[0]
            height = (self.y[self.max_row + 1] - self.y[self.min_row]) * stretch[1]
            fits = [1.0]
            if wide and width:
                fits.append(self.room[0] * wide / width)
            if tall and height:
                fits.append(self.room[1] * tall / height)
            return min(fits)
        return (setup.scale or 100) / 100

    # ── pages ────────────────────────────────────────────────────────────────
    def pages(self):
        """The printed pages, down then across, as Excel orders them; blank ones left out, as Excel does."""
        col_breaks = {b.id for b in self.ws.col_breaks.brk}
        row_breaks = {b.id for b in self.ws.row_breaks.brk}
        col_bands = self._bands(self.x, self.min_col, self.max_col, self.room[0] / self.sx, col_breaks)
        titles = self.titles
        title_height = self.y[titles[1] + 1] - self.y[titles[0]] if titles else 0.0
        row_bands = self._bands(self.y, self.min_row, self.max_row, self.room[1] / self.sy, row_breaks,
                                repeat=(titles, title_height))
        out = []
        for cols in col_bands:
            for first, last in row_bands:
                if not self._prints(first, last, cols):
                    continue
                segments = []
                top = self.margins["top"]
                if titles and first > titles[1]:
                    segments.append((titles[0], titles[1], top))
                    top += title_height * self.sy
                segments.append((first, last, top))
                out.append(Page(cols=cols, segments=segments))
        return out or [Page(cols=(self.min_col, self.max_col),
                            segments=[(self.min_row, self.max_row, self.margins["top"])])]

    def _prints(self, first, last, cols):
        if any(first <= r <= last and cols[0] <= c <= cols[1] for r, c in self.showing):
            return True
        return any(first <= p.start[2] <= last and cols[0] <= p.start[0] <= cols[1] for p in self.pictures)

    @staticmethod
    def _bands(edges, first, last, room, breaks, repeat=(None, 0.0)):
        """Runs of rows (or columns) that fit `room`, cut after each manual break."""
        titles, title_size = repeat
        bands, start, used = [], first, 0.0
        for i in range(first, last + 1):
            size = edges[i + 1] - edges[i]
            extra = title_size if titles and start > titles[1] else 0.0
            if i > start and used + size + extra > room:
                bands.append((start, i - 1))
                start, used = i, 0.0
            used += size
            if i in breaks and i < last:
                bands.append((start, i))
                start, used = i + 1, 0.0
        if start <= last:
            bands.append((start, last))
        return bands

    def _origin_x(self, cols):
        width = (self.x[cols[1] + 1] - self.x[cols[0]]) * self.sx
        left = self.margins["left"]
        if self.ws.print_options.horizontalCentered:
            left += max(0.0, (self.room[0] - width) / 2)
        return left

    # ── painting ─────────────────────────────────────────────────────────────
    def paint(self, pen, page, segment):
        first, last, top = segment
        c0, c1 = page.cols
        left = self._origin_x(page.cols)

        def px(col, inset=0.0):
            return left + (self.x[col] + inset - self.x[c0]) * self.sx

        def py(row, inset=0.0):
            return top + (self.y[row] + inset - self.y[first]) * self.sy

        def rect_of(bounds):
            r0, k0, r1, k1 = bounds
            return (px(k0), py(r0), px(k1 + 1) - px(k0), py(r1 + 1) - py(r0))

        frame = (px(c0), py(first), px(c1 + 1) - px(c0), py(last + 1) - py(first))
        cells = [cell for r in range(first, last + 1) for cell in self.rows.get(r, ())
                 if c0 <= cell.column <= c1]

        # What has text or a fill: each cell on its own, or its merged range
        # (a range that starts above or left of this page included).
        boxes = []
        for cell in cells:
            bounds = self.owner.get((cell.row, cell.column))
            if isinstance(cell, MergedCell) or (bounds and bounds[:2] != (cell.row, cell.column)):
                continue
            boxes.append((cell, bounds or (cell.row, cell.column, cell.row, cell.column)))
        for (r0, k0), bounds in self.merged.items():
            inside = first <= r0 <= last and c0 <= k0 <= c1
            overlaps = bounds[2] >= first and r0 <= last and bounds[3] >= c0 and k0 <= c1
            if overlaps and not inside:
                boxes.append((self.ws.cell(r0, k0), bounds))

        with pen.clip(*frame):
            for cell, bounds in boxes:
                fill = cell.fill
                if fill is not None and fill.fill_type == "solid":
                    color = rgb(fill.fgColor)
                    if color is not None:
                        pen.fill(*rect_of(bounds), color)
            for cell in cells:
                bounds = self.owner.get((cell.row, cell.column))
                self._borders(pen, cell, rect_of((cell.row, cell.column, cell.row, cell.column)), bounds,
                              rect_of(bounds) if bounds else None)
            for cell, bounds in boxes:
                self._text(pen, cell, bounds, rect_of(bounds), frame, page.cols, px)
            for picture in self.pictures:
                col, dx, row, dy = picture.start
                x, y = px(col, dx), py(row, dy)
                if picture.end is not None:
                    ecol, edx, erow, edy = picture.end
                    w, h = px(ecol, edx) - x, py(erow, edy) - y
                else:
                    w, h = picture.size[0] * self.scale, picture.size[1] * self.scale
                if _intersect((x, y, w, h), frame):
                    pen.image(picture.data, x, y, w, h)

    def _borders(self, pen, cell, rect, bounds, merged_rect):
        """A cell's borders. Inside a merged range only its outer edge has them, as in Excel."""
        border = cell.border
        if border is None:
            return
        x, y, w, h = rect
        r, c = cell.row, cell.column
        edges = (("top", (x, y, x + w, y), not bounds or r == bounds[0]),
                 ("bottom", (x, y + h, x + w, y + h), not bounds or r == bounds[2]),
                 ("left", (x, y, x, y + h), not bounds or c == bounds[1]),
                 ("right", (x + w, y, x + w, y + h), not bounds or c == bounds[3]))
        for side_name, (x1, y1, x2, y2), outer in edges:
            side = getattr(border, side_name)
            if outer and side is not None and side.style:
                self._rule(pen, side, x1, y1, x2, y2)
        diagonal = border.diagonal
        if (border.diagonalUp or border.diagonalDown) and diagonal is not None and diagonal.style:
            if bounds and bounds[:2] != (r, c):
                return                            # a merged range's diagonal goes with its top-left cell
            dx, dy, dw, dh = merged_rect or rect
            if border.diagonalUp:
                self._rule(pen, diagonal, dx, dy + dh, dx + dw, dy)
            if border.diagonalDown:
                self._rule(pen, diagonal, dx, dy, dx + dw, dy + dh)

    def _rule(self, pen, side, x1, y1, x2, y2):
        width, dash = BORDERS.get(side.style, BORDERS["thin"])
        width = max(width * self.scale, 0.25)
        color = rgb(side.color, BLACK)
        if side.style == "double":
            gap = max(1.2 * self.scale, 0.6)
            dx, dy = (0, gap / 2) if y1 == y2 else (gap / 2, 0)
            pen.line(x1 - dx, y1 - dy, x2 - dx, y2 - dy, width, color)
            pen.line(x1 + dx, y1 + dy, x2 + dx, y2 + dy, width, color)
        else:
            pen.line(x1, y1, x2, y2, width, color, dash and [max(d * self.scale, 0.3) for d in dash])

    def _text(self, pen, cell, bounds, rect, frame, cols, px):
        runs = runs_of(cell, self.scale)
        if not runs:
            return
        al = cell.alignment
        rotation = int(al.textRotation or 0)
        wrap, shrink = bool(al.wrap_text), bool(al.shrink_to_fit) and not al.wrap_text
        horizontal = al.horizontal or "general"
        if horizontal == "general":
            value = cell.value
            numeric = isinstance(value, (int, float, Decimal, date, datetime, time)) and not isinstance(value, bool)
            horizontal = "center" if isinstance(value, bool) else "right" if numeric else "left"
        vertical = al.vertical or "bottom"
        pad, pad_v = PAD * self.sx, PAD_V * self.sy
        indent = float(al.indent or 0) * INDENT * self.sx if horizontal in ("left", "right", "distributed") else 0.0

        if rotation == 255:                       # letters stacked top to bottom
            stacked = []
            for text, st in runs:
                for ch in text.replace(" ", ""):
                    stacked += [("\n", st), (ch, st)] if stacked else [(ch, st)]
            runs, rotation, horizontal, wrap = stacked, 0, "center", True

        if rotation:
            self._rotated(pen, runs, rotation, rect, frame, horizontal, vertical, wrap)
            return

        x, y, w, h = rect
        room = w - 2 * pad - indent
        lines = lines_of(runs, max(room, 1.0) if wrap else None)
        widest = max(line.width for line in lines)
        if shrink and widest > room > 0:
            lines = lines_of([(t, st.scaled(room / widest)) for t, st in runs])
            widest = max(line.width for line in lines)

        clip = rect
        merged = bounds[0] != bounds[2] or bounds[1] != bounds[3]
        if not wrap and not shrink and not merged and widest > room and isinstance(cell.value, (str, CellRichText)):
            clip = self._overflow(cell, rect, widest - room, horizontal, cols, px)

        block = sum(line.height for line in lines)
        if vertical == "top":
            top = y + pad_v
        elif vertical in ("center", "justify", "distributed"):
            top = y + (h - block) / 2
        else:
            top = y + h - pad_v - block
        visible = _intersect(clip, frame)
        if visible is None:
            return
        with pen.clip(*visible):
            for line in lines:
                if horizontal == "right":
                    lx = x + w - pad - indent - line.width
                elif horizontal in ("center", "centerContinuous", "distributed", "fill") or (
                        horizontal == "justify" and len(lines) == 1):
                    lx = x + (w - line.width) / 2
                else:
                    lx = x + pad + indent
                baseline = top + line.ascent
                for text, style in line.runs:
                    pen.text(lx, baseline - style.rise, text, style)
                    lx += style.width(text)
                top += line.height

    def _overflow(self, cell, rect, extra, horizontal, cols, px):
        """
        Where a too-wide text may print: on across empty neighbours, the way
        Excel lets a label run past its cell (right for left-aligned text,
        left for right-aligned, both ways for centred).
        """
        x, y, w, h = rect
        row, col = cell.row, cell.column
        filled = {c.column for c in self.rows.get(row, ()) if c.value not in (None, "")}

        def free(c):
            return cols[0] <= c <= cols[1] and c not in filled and (row, c) not in self.owner

        need_right = extra if horizontal in ("left", "justify") else extra / 2 if horizontal == "center" else 0
        need_left = extra if horizontal == "right" else extra / 2 if horizontal == "center" else 0
        right, c = x + w, col + 1
        while need_right > 0 and free(c):
            step = px(c + 1) - px(c)
            right, need_right, c = right + step, need_right - step, c + 1
        left, c = x, col - 1
        while need_left > 0 and free(c):
            step = px(c + 1) - px(c)
            left, need_left, c = left - step, need_left - step, c - 1
        return (left, y, right - left, h)

    def _rotated(self, pen, runs, rotation, rect, frame, horizontal, vertical, wrap):
        """
        Text turned on its side. Excel's rotation 1-90 turns it
        anticlockwise, 91-180 clockwise; along the turned text the cell's
        vertical alignment places it, across it the horizontal one.
        """
        x, y, w, h = rect
        angle = rotation if rotation <= 90 else -(rotation - 90)
        upward = angle > 0
        length, across = h, w                     # the text runs along the cell's height
        pad = PAD * self.sy
        lines = lines_of(runs, max(length - 2 * pad, 1.0) if wrap else None)
        block = sum(line.height for line in lines)
        if upward:
            along = {"bottom": "start", "top": "end"}.get(vertical, "middle")
            side = {"left": "top", "right": "bottom"}.get(horizontal, "middle")
        else:
            along = {"top": "start", "bottom": "end"}.get(vertical, "middle")
            side = {"right": "top", "left": "bottom"}.get(horizontal, "middle")
        across_pad = PAD * self.sx
        offset = across_pad if side == "top" else across - across_pad - block if side == "bottom" else (across - block) / 2
        visible = _intersect(rect, frame)
        if visible is None:
            return
        with pen.clip(*visible):
            for line in lines:
                if along == "start":
                    start = pad
                elif along == "end":
                    start = length - pad - line.width
                else:
                    start = (length - line.width) / 2
                depth = offset + line.ascent       # the baseline, measured across the cell
                for text, style in line.runs:
                    if upward:
                        ox, oy = x + depth - style.rise, y + h - start
                    else:
                        ox, oy = x + w - depth + style.rise, y + start
                    pen.text(ox, oy, text, style, angle=angle)
                    start += style.width(text)
                offset += line.height

    def _pictures(self):
        for img in getattr(self.ws, "_images", ()):
            anchor = img.anchor
            try:
                data = img._data()
                if isinstance(anchor, str):
                    row, col = coordinate_to_tuple(anchor)
                    yield Picture(data, (col, 0.0, row, 0.0), size=(img.width * 0.75, img.height * 0.75))
                    continue
                f = anchor._from
                start = (f.col + 1, f.colOff / EMU_PER_POINT, f.row + 1, f.rowOff / EMU_PER_POINT)
                end = getattr(anchor, "to", None)
                if end is not None:
                    yield Picture(data, start, end=(end.col + 1, end.colOff / EMU_PER_POINT,
                                                    end.row + 1, end.rowOff / EMU_PER_POINT))
                else:
                    yield Picture(data, start, size=(anchor.ext.width / EMU_PER_POINT,
                                                     anchor.ext.height / EMU_PER_POINT))
            except (AttributeError, TypeError, ValueError, OSError):
                continue                          # a picture openpyxl couldn't read is left off

    # ── header and footer ────────────────────────────────────────────────────
    def header_footer(self, pen, number, count):
        ws = self.ws
        scale = self.scale if getattr(ws.HeaderFooter, "scaleWithDoc", None) in (None, True) else 1.0
        left, right = self.margins["left"], self.paper[0] - self.margins["right"]
        for item, header in ((ws.oddHeader, True), (ws.oddFooter, False)):
            for part, align in ((item.left, "left"), (item.center, "center"), (item.right, "right")):
                if not part.text:
                    continue
                runs = _header_runs(part.text, part.font, part.size, number, count, ws.title, scale)
                if not runs:
                    continue
                line = _line(runs, runs[0][1])
                if header:
                    baseline = self.margins["header"] + line.ascent
                else:
                    descent = max(st.descent * st.size for _, st in runs)
                    baseline = self.paper[1] - self.margins["footer"] - descent
                lx = left if align == "left" else right - line.width if align == "right" else (left + right - line.width) / 2
                for text, style in line.runs:
                    pen.text(lx, baseline - style.rise, text, style)
                    lx += style.width(text)


def _header_runs(text, font, size, number, count, sheet, scale):
    """A header or footer section's codes (&P, &N, &B, ...) as runs of text."""
    name, _, kind = (font or "Calibri,Regular").partition(",")
    bold, italic, underline = "bold" in kind.lower(), "italic" in kind.lower(), False
    size = float(size or 11)
    runs = []

    def add(chunk):
        runs.append((printable(chunk), style_of(name, size, bold, italic, underline, None, None, scale)))

    for token in re.findall(r'&&|&"[^"]*"|&\d+|&K[0-9A-Fa-f]{6}|&.|[^&]+|&', text):
        if token == "&&":
            add("&")
        elif token.startswith('&"'):
            name, _, kind = token[2:-1].partition(",")
            bold, italic = "bold" in kind.lower(), "italic" in kind.lower()
        elif re.fullmatch(r"&\d+", token):
            size = float(token[1:])
        elif token == "&P":
            add(str(number))
        elif token == "&N":
            add(str(count))
        elif token == "&D":
            add(date.today().strftime("%m/%d/%Y"))
        elif token == "&T":
            add(datetime.now().strftime("%I:%M %p"))
        elif token in ("&A", "&F"):
            add(sheet)
        elif token == "&B":
            bold = not bold
        elif token == "&I":
            italic = not italic
        elif token == "&U":
            underline = not underline
        elif token.startswith("&") and len(token) > 1:
            continue                              # a colour or a code with nothing to print
        else:
            add(token)
    return [(t, st) for t, st in runs if t]


def _intersect(a, b):
    x1, y1 = max(a[0], b[0]), max(a[1], b[1])
    x2, y2 = min(a[0] + a[2], b[0] + b[2]), min(a[1] + a[3], b[1] + b[3])
    if x2 <= x1 or y2 <= y1:
        return None
    return (x1, y1, x2 - x1, y2 - y1)


# ─── The PDF itself ─────────────────────────────────────────────────────────

class Pen:
    """Draws on reportlab's canvas in points measured from the page's top left."""

    def __init__(self, title=""):
        self.buffer = BytesIO()
        self.canvas = Canvas(self.buffer, pageCompression=1)
        self.canvas.setTitle(title)
        self.canvas.setCreator("SLIS")
        self.height = 0.0
        self.started = False
        self.readers = {}

    def page(self, width, height):
        if self.started:
            self.canvas.showPage()
        self.canvas.setPageSize((width, height))
        self.height = height
        self.started = True

    def fill(self, x, y, w, h, color):
        c = self.canvas
        c.setFillColorRGB(*color)
        c.rect(x, self.height - y - h, w, h, stroke=0, fill=1)

    def line(self, x1, y1, x2, y2, width, color, dash=None):
        c = self.canvas
        c.setLineWidth(width)
        c.setStrokeColorRGB(*color)
        c.setDash(dash or [])
        c.line(x1, self.height - y1, x2, self.height - y2)

    def text(self, x, y, text, style, angle=0):
        """`text` with its baseline starting at (x, y), turned `angle` degrees anticlockwise."""
        for chunk in re.split(f"([{''.join(SHAPES)}])", text):
            if not chunk:
                continue
            if chunk in SHAPES:
                self._shape(x, y, chunk, style, angle)
            else:
                self._words(x, y, chunk, style, angle)
            step = style.width(chunk)
            x, y = x + step * cos(radians(angle)), y - step * sin(radians(angle))

    def _words(self, x, y, text, style, angle):
        c = self.canvas
        t = c.beginText()
        t.setFont(style.font, style.size)
        t.setFillColorRGB(*style.color)
        if angle:
            a = radians(angle)
            t.setTextTransform(cos(a), sin(a), -sin(a), cos(a), x, self.height - y)
        else:
            t.setTextOrigin(x, self.height - y)
        if style.hscale != 100:
            t.setHorizScale(style.hscale)
        t.textOut(text)
        c.drawText(t)
        if style.underline and not angle:
            drop = style.size * 0.1
            self.line(x, y + drop, x + style.width(text), y + drop, max(style.size * 0.05, 0.3), style.color)

    def _shape(self, x, y, mark, style, angle):
        side = SHAPE_EM * style.size
        a = radians(angle)
        c = self.canvas
        path = c.beginPath()
        for i, (u, v) in enumerate(SHAPES[mark]):
            # A corner, from the start of the baseline: the shape stands on
            # it about as tall as a capital, and turns as the text turns.
            dx, dy = u * side, (v - 0.95) * side
            px = x + dx * cos(a) + dy * sin(a)
            py = y - dx * sin(a) + dy * cos(a)
            (path.moveTo if i == 0 else path.lineTo)(px, self.height - py)
        path.close()
        c.setFillColorRGB(*style.color)
        c.drawPath(path, stroke=0, fill=1)

    def image(self, data, x, y, w, h):
        reader = self.readers.get(id(data))
        if reader is None:
            reader = self.readers[id(data)] = ImageReader(BytesIO(data))
        self.canvas.drawImage(reader, x, self.height - y - h, w, h, mask="auto")

    @contextmanager
    def clip(self, x, y, w, h):
        c = self.canvas
        c.saveState()
        path = c.beginPath()
        path.rect(x, self.height - y - h, w, h)
        c.clipPath(path, stroke=0, fill=0)
        try:
            yield
        finally:
            c.restoreState()

    def finish(self):
        if not self.started:
            self.page(*PAPER[1])
        self.canvas.showPage()
        self.canvas.save()
        return self.buffer.getvalue()
