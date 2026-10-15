"""
A print layout on one Excel sheet, drawn top to bottom: what the school's
SF10 Word templates look like, on the 8.5 x 13 in long bond they print on.

The sheet is a grid of 58 columns, each an eighth of an inch, so any table
of the templates can be laid on it to within a sixteenth of an inch. Every
value sits in a merged range, which leaves the registrar free to type over
it in Excel.

Each row gets a fixed height. That is what lets `Sheet.block` place page
breaks: a block (the top of the record, one year's table, the
certification) that won't fit on the current page starts the next one, so
no table is split across two pages. Long subject names shrink to fit
rather than wrap, for the same reason.
"""
import re
from contextlib import contextmanager

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties

from .workbook import add_picture, put

COLUMNS = 58                    # 7.25 in between the margins
EIGHTHS_PER_INCH = 8
EMU_PER_INCH = 914400
EMU_PER_POINT = 12700
# Excel measures a column in pixels at 96 dpi; a 12 px column is 1/8 in.
# Its width unit is the default font's digit, 7 px (Calibri 11).
COLUMN_WIDTH = 12 / 7

PAPER_FOLIO = 14                # Excel's "Folio, 8 1/2 x 13 in"
PAGE_HEIGHT = 13 * 72           # points
MARGINS = {"top": 0.45, "bottom": 0.4, "left": 0.62, "right": 0.62, "header": 0.2, "footer": 0.15}
# What fits between the margins, less a little for Excel's own rounding.
PAGE_BODY = PAGE_HEIGHT - (MARGINS["top"] + MARGINS["bottom"]) * 72 - 6

SERIF = "Times New Roman"
LETTERHEAD = "Century Schoolbook"

THIN = Side(style="thin", color="000000")
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
RULE = Border(bottom=THIN)      # a fill-in line under a value

def clean(text):
    """Text as a cell can hold it: what a paste can carry would make the save raise."""
    return ILLEGAL_CHARACTERS_RE.sub("", str(text or ""))


def run(text, size=8, bold=False, italic=False, underline=False, superscript=False, name=SERIF):
    """One run of a rich-text cell."""
    return TextBlock(InlineFont(
        rFont=name, sz=size, b=bold or None, i=italic or None,
        u="single" if underline else None, vertAlign="superscript" if superscript else None,
    ), clean(text))


def rich(*runs):
    return CellRichText(*[r for r in runs if r.text])


def ordinal_runs(text, size=8, bold=False, name=SERIF):
    """'21st Century' with the 'st' raised, as Word printed it."""
    out, at = [], 0
    for m in re.finditer(r"(\d)(st|nd|rd|th)\b", text):
        out.append(run(text[at:m.start(2)], size, bold, name=name))
        out.append(run(m.group(2), size, bold, superscript=True, name=name))
        at = m.end(2)
    out.append(run(text[at:], size, bold, name=name))
    return out


class Sheet:
    """Rows added top to bottom, values placed by column (0-57), page breaks kept between blocks."""

    def __init__(self, ws):
        self.ws = ws
        self.next_row = 1
        self.page_used = 0.0
        self.breaks = []            # the first row of each page after the first
        for i in range(1, COLUMNS + 1):
            ws.column_dimensions[get_column_letter(i)].width = COLUMN_WIDTH
        ws.sheet_view.showGridLines = False
        # Excel prints columns about 8% wider than it shows them, which left
        # the last few on pages of their own. Fitting one page wide shrinks
        # the printout to suit; our page breaks still hold.
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_setup.paperSize = PAPER_FOLIO
        ws.page_setup.orientation = "portrait"
        for side, inches in MARGINS.items():
            setattr(ws.page_margins, side, inches)
        ws.print_options.horizontalCentered = True
        # "Page 1 of 2", the numbers in bold, as the templates' footer has it.
        ws.oddFooter.right.text = "Page &B&P&B of &B&N&B"
        ws.oddFooter.right.font = f"{SERIF},Regular"
        ws.oddFooter.right.size = 8

    # ── rows ─────────────────────────────────────────────────────────────────
    def row(self, height):
        """A new row `height` points tall; returns its number."""
        number = self.next_row
        self.ws.row_dimensions[number].height = height
        self.next_row += 1
        return number

    def gap(self, height):
        self.row(height)

    def height(self, first, last):
        return sum(self.ws.row_dimensions[r].height or 15 for r in range(first, last + 1))

    @contextmanager
    def block(self):
        """
        Rows that print on one page. If the block doesn't fit under what the
        page already holds, a page break goes in before it.
        """
        first = self.next_row
        yield
        if self.next_row == first:
            return
        size = self.height(first, self.next_row - 1)
        if self.page_used and self.page_used + size > PAGE_BODY:
            self.ws.row_breaks.append(Break(id=first - 1))
            self.breaks.append(first)
            self.page_used = 0.0
        self.page_used += size

    def finish(self):
        self.ws.print_area = f"A1:{get_column_letter(COLUMNS)}{max(self.next_row - 1, 1)}"

    # ── values ───────────────────────────────────────────────────────────────
    def text(self, row, start, end, value, *, size=8, bold=False, italic=False, underline=False,
             h="left", v="center", wrap=False, shrink=False, border=None, rows=1, fmt=None,
             name=SERIF, indent=0):
        """
        `value` in columns start..end-1 of `row` (and the `rows` - 1 below),
        merged. The style goes on the first cell before the merge: openpyxl
        copies a range's outer border from it then, and only then.
        """
        ws = self.ws
        cell = ws.cell(row, start + 1)
        if value not in (None, ""):
            put(ws, cell.coordinate, value, number_format=fmt)
        cell.font = Font(name=name, size=size, bold=bold, italic=italic, underline="single" if underline else None)
        cell.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap or None,
                                   shrink_to_fit=shrink or None, indent=indent)
        if border is not None:
            cell.border = border
        if end - start > 1 or rows > 1:
            ws.merge_cells(start_row=row, start_column=start + 1, end_row=row + rows - 1, end_column=end)
        return cell

    def picture(self, path, row, x, width, height, y=0.0):
        """
        Float the picture at `path`, `width` x `height` inches, with its top
        left `x` inches across from column 0 and `y` points down from the
        top of `row`.
        """
        while y >= (self.ws.row_dimensions[row].height or 15) and row < self.next_row - 1:
            y -= self.ws.row_dimensions[row].height or 15
            row += 1
        col = int(x * EIGHTHS_PER_INCH)
        add_picture(
            self.ws, path, col=col, row=row - 1,
            col_off=round((x - col / EIGHTHS_PER_INCH) * EMU_PER_INCH),
            row_off=round(y * EMU_PER_POINT),
            cx=round(width * EMU_PER_INCH), cy=round(height * EMU_PER_INCH),
        )


def edges(widths):
    """Column widths (in eighths) to where each starts: (5, 24) -> [0, 5, 29]."""
    out = [0]
    for w in widths:
        out.append(out[-1] + w)
    return out
