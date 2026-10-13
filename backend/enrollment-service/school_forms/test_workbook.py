"""The openpyxl helpers: safe writes, merged ranges, and growing a learner list."""
import io

import openpyxl
import pytest

from school_forms import workbook
from school_forms.workbook import ensure_merged, insert_rows, merge_rows, put, shift_ref


def _sheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Header"
    for row in (2, 3):                       # a two-row learner block
        ws.merge_cells(f"B{row}:C{row}")
        ws.row_dimensions[row].height = 30
        ws[f"A{row}"].font = openpyxl.styles.Font(bold=True)
    ws["A4"] = "Footer"
    ws.merge_cells("A5:D5")
    ws["A5"] = "Signature"
    ws.row_dimensions[4].height = 12
    ws.print_area = "A1:D5"
    return wb, ws


def test_put_keeps_text_as_text():
    wb, ws = _sheet()
    cell = put(ws, "A2", '=HYPERLINK("http://x","y")')
    assert cell.data_type == "s"
    reloaded = openpyxl.load_workbook(io.BytesIO(workbook.to_bytes(wb)))
    assert reloaded.active["A2"].value == '=HYPERLINK("http://x","y")'
    assert reloaded.active["A2"].data_type == "s"


def test_put_strips_characters_excel_cannot_store():
    wb, ws = _sheet()
    put(ws, "A2", "Blk 1\x0b Lot 2")
    assert ws["A2"].value == "Blk 1 Lot 2"
    workbook.to_bytes(wb)  # would raise IllegalCharacterError otherwise


def test_put_writes_a_merged_range_from_any_of_its_cells():
    _, ws = _sheet()
    put(ws, "C2", "Dela Cruz")
    assert ws["B2"].value == "Dela Cruz"


def test_put_shrink_sets_upright_shrink_to_fit():
    _, ws = _sheet()
    ws["D2"].alignment = openpyxl.styles.Alignment(textRotation=90, wrap_text=True)
    cell = put(ws, "D2", 51, number_format="0", shrink=True)
    assert cell.alignment.shrink_to_fit and not cell.alignment.wrap_text
    assert cell.alignment.textRotation == 0
    assert cell.number_format == "0"


def test_insert_rows_moves_everything_below_and_styles_the_new_rows():
    _, ws = _sheet()
    insert_rows(ws, 4, 3, style_row=3)
    assert ws["A7"].value == "Footer"
    assert ws.row_dimensions[7].height == 12
    assert ws["A8"].value == "Signature"
    coords = {r.coord for r in ws.merged_cells.ranges}
    assert "A8:D8" in coords and "A5:D5" not in coords
    for row in (4, 5, 6):
        assert f"B{row}:C{row}" in coords
        assert ws.row_dimensions[row].height == 30
        assert ws[f"A{row}"].font.b
    assert ws.print_area.endswith("$A$1:$D$8")


def test_insert_rows_refuses_to_split_a_merged_range():
    _, ws = _sheet()
    ws.merge_cells("D2:D3")
    with pytest.raises(ValueError):
        insert_rows(ws, 3, 1, style_row=2)


def test_insert_rows_of_nothing_changes_nothing():
    _, ws = _sheet()
    insert_rows(ws, 4, 0, style_row=3)
    assert ws["A4"].value == "Footer"


def test_merge_rows_makes_every_row_match():
    _, ws = _sheet()
    ws.unmerge_cells("B3:C3")
    ws.merge_cells("C3:D3")                  # a template's odd row
    merge_rows(ws, range(2, 4), ("B:C",))
    coords = {r.coord for r in ws.merged_cells.ranges}
    assert {"B2:C2", "B3:C3"} <= coords and "C3:D3" not in coords


def test_ensure_merged_replaces_an_overlap_and_is_idempotent():
    _, ws = _sheet()
    ensure_merged(ws, "B2:D2")
    ensure_merged(ws, "B2:D2")
    coords = [r.coord for r in ws.merged_cells.ranges]
    assert coords.count("B2:D2") == 1 and "B2:C2" not in coords


def test_shift_ref():
    assert shift_ref("R59", 3) == "R62"
    assert shift_ref("AH66", 0) == "AH66"
