"""
SF10 in the school's own format: grade rules, the three layouts, the
certification and the page breaks.

The expected numbers are the school's: they come off its 2026 SF10
templates (SF10-ES, SF10-JHS, SF10-SHS HUMSS), so a change that makes the
export disagree with how the school computes a record fails here.
"""
from datetime import date
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from openpyxl.cell.cell import MergedCell

from school_forms import sf10
from school_forms.records import (
    IssueDetails, Learner, PermanentRecordData, PriorLevel, School, Signatories, SubjectRecord, YearRecord,
)
from school_forms.sf10_sheet import PAGE_BODY
from school_forms.workbook import to_bytes

SCHOOL = School(
    name="South Lakes Integrated School", address="607 Brgy. San Jose, San Pablo City, Laguna 4000",
    school_id="428023", division="San Pablo City", phone="+63998-979-1547", landline="(049) 560-3268",
    email="southlakesintegratedschool@yahoo.com.ph",
    signatories=Signatories(records_officer="Regina T. Molidor, LPT", elementary_principal="Maria Carolina A. Laguerta, LPT",
                            jhs_principal="Susan Charito A. Agoncillo, LPT", shs_principal="Maria Carolina A. Laguerta, LPT",
                            shs_department_head="Erwin V. Bustamante, M. Div."),
)
LEARNER = Learner(enrollment_id=0, student_id=5, lrn="109100000005", last_name="Dela Cruz", first_name="Maria",
                  middle_name="Santos", sex="female", birth_date=date(2014, 4, 30),
                  address="Barangay Del Remedio, San Pablo City, Laguna", father="Juan Dela Cruz",
                  mother="Ana Dela Cruz", student_number="2019-0042")
Q = ("1st_quarter", "2nd_quarter", "3rd_quarter", "4th_quarter")
KINDER = PriorLevel(level="KINDER", school_year="2019-2020", school=SCHOOL.name, final_rating="-",
                    remarks="Grade 1 Ready")
_order = iter(range(1, 10_000))


def subject(name, grades, remarks=()):
    return SubjectRecord(name=name, grades={p: Decimal(g) for p, g in zip(Q, grades) if g is not None},
                         recorded_remarks=tuple(remarks), order=next(_order))


def mapeh_parts(music, arts, pe, health):
    return [subject("Music", music), subject("Arts", arts), subject("PE", pe), subject("Health", health)]


def year(grade, sy, subjects, section="Laguna Lake", completed=True):
    return YearRecord(school_year=sy, grade_level=grade, section=section, periods=Q, subjects=tuple(subjects),
                      completed=completed)


def semester(grade, sy, number, subjects, strand="HUMSS", section="Kayangan Lake"):
    key = f"{number}_semester"
    records = tuple(
        SubjectRecord(name=name, grades={key: Decimal(g)}, order=next(_order),
                      group=sf10.subject_group(name, strand if tied else None))
        for name, g, tied in subjects
    )
    return YearRecord(school_year=sy, grade_level=grade, section=section, periods=(key,), subjects=records,
                      strand=strand, semester=number, completed=True)


# The school's own Grade 1 and Grade 5 from its SF10-ES template.
GRADE_1 = year("Grade 1", "2020-2021", [
    subject("English 1", [93, 95, 96, 97]), subject("Mathematics 1", [91, 91, 92, 95]),
    subject("Science 1", [90, 92, 94, 96]), subject("Filipino 1", [93, 94, 96, 96]),
    subject("Araling Panlipunan 1", [94, 95, 95, 97]), *mapeh_parts(*[[94, 95, 96, 98]] * 4),
    subject("Computer", [90, 91, 91, 91]), subject("Values", [94, 94, 94, 97]),
])
GRADE_5 = year("Grade 5", "2024-2025", [
    subject("English", [83, 87, 88, 88]), subject("Math", [87, 88, 91, 91]), subject("Science", [88, 86, 85, 86]),
    subject("Filipino", [84, 87, 87, 90]), subject("Araling Panlipunan", [84, 85, 86, 87]),
    *mapeh_parts([87, 89, 87, 90], [84, 86, 85, 89], [90, 90, 85, 90], [89, 90, 90, 92]),
    subject("Values", [85, 85, 85, 86]), subject("Computer", [91, 90, 90, 90]), subject("HELE", [85, 85, 87, 88]),
], section="Naujan Lake")
# Its Grade 9 from SF10-JHS: Environmental Science ran two quarters.
GRADE_9 = year("Grade 9", "2024-2025", [
    subject("English", [95, 96, 97, 96]), subject("Mathematics", [95, 95, 97, 97]),
    subject("Science", [95, 94, 98, 97]), subject("Filipino", [96, 96, 97, 97]),
    subject("Araling Panlipunan", [97, 97, 97, 97]),
    *mapeh_parts([96, 93, 95, 95], [93, 92, 96, 95], [96, 94, 95, 95], [94, 95, 97, 96]),
    subject("Values Education", [93, 97, 98, 98]), subject("TLE-ICT", [94, 93, 94, 99]),
    subject("TLE-HE", [96, 97, 96, 97]), subject("Environmental Science", [95, 95, None, None]),
], section="Labas Lake")


def cells(ws, row):
    """A row's filled cells, left to right: each merged range once."""
    return [c for c in ws[row] if not isinstance(c, MergedCell) and c.value not in (None, "")]


def values(ws, row):
    """What a row shows; rich text as its plain text."""
    return [c.value if isinstance(c.value, (int, float, str)) else str(c.value) for c in cells(ws, row)]


def rows_starting(ws, text):
    return [r for r in range(1, ws.max_row + 1) if values(ws, r)[:1] == [text]]


def spans(ws, row):
    """The merged ranges in `row` as (first column, last column), counted from 1, left to right."""
    return sorted((r.min_col, r.max_col) for r in ws.merged_cells.ranges if r.min_row <= row <= r.max_row)


def find(ws, text):
    """The first row with a cell that reads `text` (rich text included)."""
    for r in range(1, ws.max_row + 1):
        if any(str(v) == text for v in values(ws, r)):
            return r
    raise AssertionError(f"{text!r} not on the sheet")


def all_text(ws):
    return " ".join(str(v) for r in range(1, ws.max_row + 1) for v in values(ws, r))


def build(level, years, prior=None, issue=IssueDetails(copy_for="Pamantasan ng Lungsod ng San Pablo"),
          learner=LEARNER, school=SCHOOL):
    wb, warnings = sf10.build(PermanentRecordData(school=school, level=level, learner=learner, years=tuple(years),
                                                  prior=prior, issue=issue))
    return wb.active, warnings


# ── Grade rules ─────────────────────────────────────────────────────────────

def test_final_rating_rounds_half_up_like_deped():
    # 74, 75, 74, 75 average 74.5: the whole-number final is 75, passed.
    grades, final, remark = sf10.subject_result(subject("Math", [74, 75, 74, 75]), Q)
    assert grades == [74, 75, 74, 75] and final == 75 and remark == "passed"


def test_a_running_year_waits_for_every_quarter():
    grades, final, remark = sf10.subject_result(subject("Math", [90, 91, None, None]), Q)
    assert grades == [90, 91, None, None] and final is None and remark is None


def test_a_completed_year_rates_a_subject_on_the_quarters_it_ran():
    grades, final, remark = sf10.subject_result(subject("Environmental Science", [95, 95, None, None]), Q, True)
    assert grades == [95, 95, "-", "-"] and final == 95 and remark == "passed"


def test_a_teachers_dropped_mark_wins_over_the_numbers():
    _, final, remark = sf10.subject_result(subject("Math", [90] * 4, remarks=["dropped"]), Q)
    assert final == 90 and remark == "dropped"


def test_mapeh_is_its_parts_unrounded_mean_like_the_schools_records():
    result = sf10.year_result(GRADE_5)
    mapeh = next(l for l in result.lines if l.label == "MAPEH")
    # Quarters print rounded; the final averages 87.5, 88.75, 86.75, 90.25.
    assert mapeh.grades == [88, 89, 87, 90] and mapeh.final == 88
    assert [l.label for l in result.lines if l.part] == ["Music", "Arts", "P.E.", "Health"]
    assert (result.period_averages, result.average) == ([86, 87, 87, 88], 87)


def test_mapeh_without_its_parts_prints_them_blank_and_says_so():
    result = sf10.year_result(year("Grade 3", "2022-2023", [subject("MAPEH 3", [91, 91, 91, 93])]))
    assert [(l.label, l.grades) for l in result.lines] == [
        ("MAPEH", [91, 91, 91, 93]), ("Music", [None] * 4), ("Arts", [None] * 4),
        ("P.E.", [None] * 4), ("Health", [None] * 4),
    ]
    assert "aren't recorded as separate subjects" in result.notes[0]


def test_a_lone_health_subject_is_not_folded_into_mapeh():
    result = sf10.year_result(year("Grade 7", "2022-2023", [subject("Health", [90] * 4)]))
    assert [l.label for l in result.lines] == ["Health"] and result.average == 90


def test_the_schools_order_of_learning_areas():
    names = ["Intro to Gen Math", "Computer", "MAPEH", "Values Education", "Filipino 7", "Science 7",
             "Mathematics 7", "English 7", "Araling Panlipunan 7", "Journalism"]
    result = sf10.year_result(year("Grade 7", "2022-2023", [subject(n, [90] * 4) for n in names]))
    labels = [l.label for l in result.lines if not l.part]
    # Electives follow in the order they were entered, wherever "Math" sits in the name.
    assert labels == ["English", "Mathematics", "Science", "Filipino", "Araling Panlipunan", "MAPEH",
                      "Values Education", "Computer", "Intro to Gen Math", "Journalism"]


def test_quarter_averages_count_the_subjects_taken_that_quarter():
    result = sf10.year_result(GRADE_9)
    assert (result.period_averages, result.average) == ([95, 95, 97, 97], 96)
    assert sf10.year_result(GRADE_1).period_averages == [92, 93, 94, 96]


def test_a_year_in_progress_has_no_general_average():
    result = sf10.year_result(year("Grade 4", "2026-2027", [subject("English", [90, 91, None, None])],
                                   completed=False))
    assert result.period_averages == [90, 91, None, None] and result.average is None and not result.passed


def test_one_or_two_failed_areas_mean_remedial_three_retained():
    def outcome(fails):
        subjects = [subject(f"Subject {i}", [70] * 4 if i < fails else [90] * 4) for i in range(4)]
        return sf10._outcome(sf10.year_result(year("Grade 4", "2023-2024", subjects)), False)
    assert (outcome(0), outcome(1), outcome(2), outcome(3)) == ("Passed", "Remedial", "Remedial", "Retained")


@pytest.mark.parametrize("name, strand, stored, expected", [
    ("Oral Communication in Context", None, None, "core"),
    ("Physical Education and Health 1", None, None, "core"),
    ("Empowerment Technologies", None, None, "applied"),
    ("Research in Daily Life 1", "HUMSS", None, "applied"),      # applied even when tied to a strand
    ("Research Project/3Is", None, None, "applied"),
    ("Pagsulat sa Filipino sa Piling Larangan (Akademik)", None, None, "applied"),
    ("Pre-Calculus", "STEM", None, "specialized"),
    ("Culminating Activity", None, None, "specialized"),
    ("Research/Capstone Project", "STEM", None, "specialized"),
    ("Pre-Calculus", "STEM", "core", "core"),                    # the registrar's choice wins
])
def test_senior_high_subject_groups(name, strand, stored, expected):
    assert sf10.subject_group(name, strand, stored) == expected


def test_senior_high_hours_and_strand_names():
    assert sf10.subject_hours("Physical Education and Health 4") == 20
    assert sf10.subject_hours("General Mathematics") == 80
    assert sf10.strand_name("stem") == "Science, Technology, Engineering, and Mathematics (STEM)"
    assert sf10.strand_name("Robotics") == "Robotics"


# ── The elementary record ───────────────────────────────────────────────────

def test_elementary_record_matches_the_schools_template():
    ws, warnings = build("elementary", [GRADE_1], prior=KINDER)
    assert warnings == []
    assert ws.title == "SF10-ES"
    assert values(ws, 1) == ["SLIS SF10-ES – Learner Permanent Record for Elementary School", "Student No: 2019-0042"]
    assert values(ws, 2) == ["School ID: 428023", "LRN: 109100000005"]
    text = all_text(ws)
    assert "ELEMENTARY SCHOOL PERMANENT RECORD" in text and "Mobile: +63998-979-1547 | Landline: (049) 560-3268" in text
    name_row = find(ws, "DELA CRUZ, MARIA SANTOS")
    assert values(ws, name_row) == ["Name", "DELA CRUZ, MARIA SANTOS", "Division", "SAN PABLO CITY", "School", "SLIS"]
    assert values(ws, name_row + 1) == ["Sex", "F", "Date of Birth", "APRIL 30, 2014"]
    assert values(ws, name_row + 2) == ["Parents", "JUAN DELA CRUZ / ANA DELA CRUZ", "Address",
                                        "BARANGAY DEL REMEDIO, SAN PABLO CITY, LAGUNA"]
    assert values(ws, find(ws, "KINDER")) == ["2019-2020", "SOUTH LAKES INTEGRATED SCHOOL", "KINDER", "-", "Grade 1 Ready"]

    (top,) = rows_starting(ws, "Grade & Section")
    assert values(ws, top) == ["Grade & Section", "1 – LAGUNA LAKE", "School ID: 428023", "SY 2020-2021", "Promoted to:"]
    assert values(ws, top + 1) == ["School", "SOUTH LAKES INTEGRATED SCHOOL", "GRADE 2"]
    assert values(ws, find(ws, "English")) == ["English", 93, 95, 96, 97, 95, "Passed"]
    assert values(ws, find(ws, "MAPEH")) == ["MAPEH", 94, 95, 96, 98, 96, "Passed"]
    assert values(ws, find(ws, "Music")) == ["Music", 94, 95, 96, 98]
    assert values(ws, find(ws, "General Average")) == ["General Average", 92, 93, 94, 96, 94, "Passed"]
    assert ws.cell(find(ws, "English"), 1).font.bold
    final = cells(ws, find(ws, "General Average"))[5]
    assert (final.value, final.number_format, final.font.bold) == (94, "0.00", True)


def test_the_certification():
    ws, _ = build("elementary", [GRADE_1], prior=KINDER)
    text = all_text(ws)
    assert ("I hereby certify that this is a true record of MARIA SANTOS DELA CRUZ. I further certify that based on "
            "the information contained herein, she is eligible for admission to Grade 2.") in text
    assert ("issued upon the request of the above-named person for TRANSFER CREDENTIAL PURPOSES to "
            "Pamantasan ng Lungsod ng San Pablo only.") in text
    assert "Elementary School Principal" in text
    names = find(ws, "Regina T. Molidor, LPT")
    assert values(ws, names) == ["Regina T. Molidor, LPT", "Maria Carolina A. Laguerta, LPT"]
    assert values(ws, names + 1) == ["Records Officer", "Elementary School Principal"]
    assert "Copy for Pamantasan ng Lungsod ng San Pablo" in text


def test_with_no_names_set_the_signature_lines_print_blank():
    ws, _ = build("elementary", [GRADE_1], prior=KINDER, school=School(name=SCHOOL.name))
    titles = find(ws, "Records Officer")
    assert values(ws, titles) == ["Records Officer", "Elementary School Principal"]
    assert values(ws, titles - 1) == []                            # the names, left to write in
    lines = spans(ws, titles - 2)
    assert len(lines) == 2 and all(ws.cell(titles - 2, first).border.bottom.style == "thin" for first, _ in lines)


def test_a_copy_with_nothing_typed_prints_lines_to_fill_in():
    ws, _ = build("elementary", [GRADE_1], prior=KINDER, issue=IssueDetails(),
                  learner=Learner(**{**LEARNER.__dict__, "sex": "male"}))
    text = all_text(ws)
    assert "he is eligible for admission to Grade 2" in text
    assert f"TRANSFER CREDENTIAL PURPOSES to {sf10.BLANK} only." in text and f"Copy for {sf10.BLANK}" in text


def test_kindergarten_alone_makes_an_elementary_record():
    ws, warnings = build("elementary", [], prior=KINDER)
    text = all_text(ws)
    assert "KINDERGARTEN PROGRESS" in text and "ELEMENTARY PROGRESS" not in text
    assert "she is eligible for admission to Grade 1." in text
    assert warnings == []


def test_a_level_not_taken_here_leaves_its_row_to_fill_in():
    ws, warnings = build("junior_highschool", [GRADE_9], prior=PriorLevel(level="GRADE 6"))
    assert values(ws, find(ws, "GRADE 6")) == ["GRADE 6"]
    assert any("Grade 6 wasn't taken here" in w for w in warnings)


def test_missing_signatory_names_are_warned_about():
    _, warnings = build("elementary", [GRADE_1], prior=KINDER, school=School(name=SCHOOL.name))
    assert [w for w in warnings if w.startswith("No name is set")] == [
        "No name is set for the Elementary School Principal under Settings › School Forms, "
        "so that signature line prints blank.",
        "No name is set for the Records Officer under Settings › School Forms, so that signature line prints blank.",
    ]


# ── Junior high ─────────────────────────────────────────────────────────────

def test_junior_high_record():
    prior = PriorLevel(level="GRADE 6", school_year="2021-2022", school=SCHOOL.name, final_rating="95",
                       remarks="PROMOTED")
    ws, warnings = build("junior_highschool", [GRADE_9], prior=prior)
    assert warnings == []
    assert values(ws, find(ws, "GRADE 6")) == ["2021-2022", "SOUTH LAKES INTEGRATED SCHOOL", "GRADE 6", 95, "PROMOTED"]
    (top,) = rows_starting(ws, "Grade/Year Level")
    assert values(ws, top)[:2] == ["Grade/Year Level", "9 – LABAS LAKE"]
    assert values(ws, top + 1)[-1] == "GRADE 10"
    assert values(ws, find(ws, "Environmental Science")) == ["Environmental Science", 95, 95, "-", "-", 95, "Passed"]
    assert values(ws, find(ws, "General Average")) == ["General Average", 95, 95, 97, 97, 96, "Passed"]
    text = all_text(ws)
    assert "eligible for admission to Grade 10." in text
    assert "Junior High School Principal/Administrator" in text


# ── Senior high ─────────────────────────────────────────────────────────────

HUMSS = [
    semester("Grade 11", "2024-2025", "1st", [
        ("Oral Communication in Context", 93, 0), ("General Mathematics", 91, 0),
        ("Physical Education and Health 1", 94, 0), ("Understanding Culture, Society, and Politics", 96, 0),
        ("Earth and Life Science", 91, 0), ("21st Century Literature from the Philippines and the World", 91, 0),
        ("Komunikasyon at Pananaliksik sa Wika at Kulturang Pilipino", 94, 0),
        ("Research in Daily Life 1", 92, 0), ("Empowerment Technologies", 89, 0)]),
    semester("Grade 11", "2024-2025", "2nd", [
        ("Reading and Writing Skills", 94, 0), ("Statistics and Probability", 93, 0),
        ("English for Academic and Professional Purposes", 96, 0),
        ("Philippine Politics and Governance", 94, 1)]),
    semester("Grade 12", "2025-2026", "1st", [("Media and Information Literacy", 93, 0)], section="Buhi Lake"),
    semester("Grade 12", "2025-2026", "2nd", [
        ("Physical Science", 91, 0), ("Physical Education and Health 4", 96, 0),
        ("Culminating Activity", 95, 1)], section="Buhi Lake"),
]
GRADUATE = IssueDetails(copy_for="Adamson University", graduation_date=date(2026, 3, 31), special_order="051")


def test_senior_high_record_matches_the_humss_template():
    ws, warnings = build("senior_highschool", HUMSS,
                         prior=PriorLevel(level="GRADE 10", school_year="2023-2024", school=SCHOOL.name,
                                          final_rating="89", remarks="PASSED"),
                         issue=GRADUATE)
    assert warnings == ["SLIS records one grade per senior high semester, so the quarter columns are blank "
                        "and Final Rating is the semester grade."]
    name_row = find(ws, "DELA CRUZ, MARIA SANTOS")
    assert values(ws, name_row + 1) == ["Sex", "FEMALE", "Date of Birth", "APRIL 30, 2014"]
    top = rows_starting(ws, "Grade/Year Level")[0]
    assert values(ws, top + 2) == ["FIRST SEMESTER", "Humanities and Social Sciences (HUMSS)"]
    assert values(ws, top + 3) == ["SUBJECTS", "1st Quarter", "2nd Quarter", "Final Rating", "No. of Hours", "Action Taken"]
    # Core, PE and Health last; then Contextualized; quarters blank.
    first_semester = [values(ws, r)[0] for r in range(top + 4, top + 16)]
    assert first_semester == [
        "Core Subjects", "Oral Communication in Context", "General Mathematics",
        "Understanding Culture, Society, and Politics", "Earth and Life Science",
        "21st Century Literature from the Philippines and the World",
        "Komunikasyon at Pananaliksik sa Wika at Kulturang Pilipino", "Physical Education and Health 1",
        "Contextualized Subjects", "Research in Daily Life 1", "Empowerment Technologies",
        "General Average for the Semester",
    ]
    assert values(ws, find(ws, "Oral Communication in Context")) == ["Oral Communication in Context", 93, 80, "Passed"]
    assert values(ws, find(ws, "Physical Education and Health 1"))[2] == 20
    assert values(ws, top + 15) == ["General Average for the Semester", 92, 660, "Passed"]
    assert cells(ws, top + 15)[2].number_format == '0" hours"'

    second = rows_starting(ws, "Grade/Year Level")[1]
    assert values(ws, second + 3)[1:3] == ["3rd Quarter", "4th Quarter"]
    assert "Specialization Subjects" in [values(ws, r)[0] for r in range(second + 4, second + 12)]
    promoted = [values(ws, r) for r in range(1, ws.max_row + 1) if values(ws, r)[:1] == ["Promoted to:"]]
    assert promoted == [["Promoted to:", "GRADE 12"], ["Promoted to:", "COLLEGE"]]

    text = all_text(ws)
    assert "Date of Graduation: March 31, 2026" in text and "Special Order (A) No. 051, s. 2026" in text
    assert "she is eligible for admission to College." in text
    assert "Erwin V. Bustamante, M. Div." in text and "Senior High School Department Head" in text
    assert "Records Officer/Registrar" in text and "Copy for Adamson University" in text


def test_a_grade_still_in_its_first_semester_promotes_nobody():
    ws, _ = build("senior_highschool", HUMSS[:1], issue=GRADUATE)
    text = all_text(ws)
    assert "Promoted to:" not in text and "Date of Graduation" not in text
    assert f"eligible for admission to {sf10.BLANK}." in text


@pytest.mark.parametrize("level, years, blocks", [
    ("elementary", [GRADE_1], [(1, 24), (35, 58)]),
    ("junior_highschool", [GRADE_9], [(1, 24), (35, 58)]),
    ("senior_highschool", HUMSS, [(1, 18), (21, 38), (41, 58)]),
])
def test_the_signatures_stand_side_by_side(level, years, blocks):
    ws, _ = build(level, years)
    labels = find(ws, "Prepared by:")
    assert values(ws, labels) == ["Prepared by:", "Certified True and Correct:"]
    assert [first for first, _ in spans(ws, labels)] == [blocks[0][0], blocks[1][0]]
    names = find(ws, "Regina T. Molidor, LPT")
    # One row each for the lines, the names and the titles, every signature in the same columns.
    assert spans(ws, names - 1) == spans(ws, names) == spans(ws, names + 1) == blocks
    assert all(ws.cell(names - 1, first).border.bottom.style == "thin" for first, _ in blocks)
    assert values(ws, names + 1)[0] == sf10.LAYOUTS[level].records_officer_title
    # The date goes under the records officer's signature.
    assert spans(ws, find(ws, "Date:")) == [(1, 4), (5, blocks[0][1])]


# ── The sheet ───────────────────────────────────────────────────────────────

def test_page_setup_is_the_schools_long_bond():
    ws, _ = build("elementary", [GRADE_1], prior=KINDER)
    assert ws.page_setup.paperSize == 14 and ws.page_setup.orientation == "portrait"
    # One page wide however much wider Excel prints the columns; our page breaks still decide the rest.
    assert ws.sheet_properties.pageSetUpPr.fitToPage
    assert (ws.page_setup.fitToWidth, ws.page_setup.fitToHeight) == (1, 0)
    assert ws.sheet_view.showGridLines is False
    assert "&P" in ws.oddFooter.right.text and "&N" in ws.oddFooter.right.text
    assert len(ws._images) == 2           # the crest and the photo box


def _page_starts(ws):
    return [1] + [b.id + 1 for b in ws.row_breaks.brk] + [ws.max_row + 1]


@pytest.mark.parametrize("level, years", [
    ("elementary", [GRADE_1, GRADE_5] * 3),
    ("senior_highschool", HUMSS * 2),
])
def test_pages_break_between_tables_never_inside_one(level, years):
    ws, _ = build(level, years, prior=KINDER if level == "elementary" else None)
    starts = _page_starts(ws)
    assert len(starts) > 2                                    # it took more than one page
    label = "Grade & Section" if level == "elementary" else "Grade/Year Level"
    blocks = set(rows_starting(ws, label)) | {find(ws, "C E R T I F I C A T I O N") - 1}
    assert set(starts[1:-1]) <= blocks
    for first, end in zip(starts, starts[1:]):
        assert sum(ws.row_dimensions[r].height for r in range(first, end)) <= PAGE_BODY


def test_a_long_subject_name_wraps_onto_a_second_line():
    long_name = "Edukasyong Pantahanan at Pangkabuhayan"
    ws, _ = build("elementary", [year("Grade 4", "2023-2024", [subject(long_name, [90] * 4),
                                                              subject("English", [90] * 4)])], prior=KINDER)
    long_row, short_row = find(ws, long_name), find(ws, "English")
    assert ws.row_dimensions[long_row].height == 2 * ws.row_dimensions[short_row].height
    assert ws.cell(long_row, 1).alignment.wrap_text and not ws.cell(short_row, 1).alignment.wrap_text


def test_text_that_looks_like_a_formula_stays_text():
    sneaky = year("Grade 4", "2023-2024", [subject('=HYPERLINK("http://x","Math")', [90] * 4)])
    wb, _ = sf10.build(PermanentRecordData(school=SCHOOL, level="elementary", learner=LEARNER,
                                           years=(sneaky,), prior=KINDER))
    ws = openpyxl.load_workbook(BytesIO(to_bytes(wb))).active
    cell = ws.cell(find(ws, '=HYPERLINK("http://x","Math")'), 1)
    assert cell.data_type == "s"
